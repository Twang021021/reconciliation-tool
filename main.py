"""Spreadsheet reconciliation tool.

Compares two versions of the same tabular data (CSV or XLSX), pairs rows by
a key column, and sorts every row into one of four buckets: clean matches,
field-level mismatches, "only in A", and "only in B". Duplicate keys are
flagged separately since they break the 1:1 pairing this tool relies on.
"""

import argparse
import os
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ReconciliationError(Exception):
    """A problem with the input files or config that the user needs to fix,
    as opposed to an unexpected internal error."""

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
CONFIG = {
    "file_a": "sample_data/file_a.csv",
    "file_b": "sample_data/file_b.csv",
    "key_column": "id",
    "numeric_tolerance": 0.01,
    # Maps a column name in file B to the equivalent column name in file A,
    # for cases where the same field is labeled differently in each file.
    "column_mapping": {
        "full_name": "name",
    },
    "output_dir": "output",
    "output_excel_path": "output/reconciliation_report.xlsx",
}


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------
def load_file(path):
    if not os.path.isfile(path):
        raise ReconciliationError(f"File not found: {path}")

    ext = Path(path).suffix.lower()
    try:
        if ext == ".csv":
            df = pd.read_csv(path, dtype=str, keep_default_na=False)
        elif ext in (".xlsx", ".xls"):
            df = pd.read_excel(path, dtype=str)
        else:
            raise ReconciliationError(
                f"Unsupported file type '{ext}': {path} (expected .csv, .xlsx, or .xls)"
            )
    except pd.errors.EmptyDataError:
        raise ReconciliationError(f"File is empty: {path}")

    if df.shape[1] == 0:
        raise ReconciliationError(f"File has no columns: {path}")

    return df


# ---------------------------------------------------------------------------
# Value normalization / comparison
# ---------------------------------------------------------------------------
_CURRENCY_PREFIX = re.compile(r"^[$€£]")


def is_blank(value):
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def try_parse_number(value):
    """Best-effort numeric parse, tolerant of currency/thousands formatting."""
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = _CURRENCY_PREFIX.sub("", value.strip()).replace(",", "")
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def normalize_key(value):
    """Stable string form of a key value, so '1', ' 1 ', and 1.0 all match."""
    if is_blank(value):
        return None
    num = try_parse_number(value)
    if num is not None:
        return str(int(num)) if num == int(num) else str(num)
    return str(value).strip().lower()


def values_equal(a, b, tolerance):
    a_blank, b_blank = is_blank(a), is_blank(b)
    if a_blank and b_blank:
        return True
    if a_blank != b_blank:
        return False

    num_a, num_b = try_parse_number(a), try_parse_number(b)
    if num_a is not None and num_b is not None:
        return abs(num_a - num_b) <= tolerance

    return str(a).strip().lower() == str(b).strip().lower()


# ---------------------------------------------------------------------------
# Reconciliation
# ---------------------------------------------------------------------------
def _extract_side(df, suffix, original_columns, key):
    """Pull the suffixed columns from a one-sided merge result back into
    plain original column names, dropping merge bookkeeping columns."""
    cols = [c for c in original_columns if c != "__key__"]
    out = pd.DataFrame({col: df.get(f"{col}{suffix}", df.get(col)) for col in cols})
    return out.reset_index(drop=True)


def reconcile(config):
    df_a = load_file(config["file_a"])
    df_b = load_file(config["file_b"]).rename(columns=config["column_mapping"])

    key = config["key_column"]
    tolerance = config["numeric_tolerance"]

    if key not in df_a.columns:
        raise ReconciliationError(
            f"Key column '{key}' not found in file A. Available columns: {list(df_a.columns)}"
        )
    if key not in df_b.columns:
        raise ReconciliationError(
            f"Key column '{key}' not found in file B (after applying column_mapping). "
            f"Available columns: {list(df_b.columns)}"
        )

    df_a["__key__"] = df_a[key].map(normalize_key)
    df_b["__key__"] = df_b[key].map(normalize_key)

    dup_a_mask = df_a["__key__"].duplicated(keep=False) & df_a["__key__"].notna()
    dup_b_mask = df_b["__key__"].duplicated(keep=False) & df_b["__key__"].notna()

    duplicates_a = df_a[dup_a_mask].drop(columns="__key__")
    duplicates_b = df_b[dup_b_mask].drop(columns="__key__")

    clean_a = df_a[~dup_a_mask]
    clean_b = df_b[~dup_b_mask]

    common_cols = [c for c in df_a.columns if c in df_b.columns and c not in (key, "__key__")]
    only_in_a_cols = [c for c in df_a.columns if c not in df_b.columns and c != "__key__"]
    only_in_b_cols = [c for c in df_b.columns if c not in df_a.columns and c != "__key__"]

    merged = pd.merge(
        clean_a, clean_b, on="__key__", how="outer",
        suffixes=("_a", "_b"), indicator=True,
    )

    only_in_a = _extract_side(merged[merged["_merge"] == "left_only"], "_a", df_a.columns, key)
    only_in_b = _extract_side(merged[merged["_merge"] == "right_only"], "_b", df_b.columns, key)
    both = merged[merged["_merge"] == "both"]

    clean_rows = []
    mismatch_rows = []

    for _, row in both.iterrows():
        key_val = row.get(f"{key}_a", row.get(key))
        mismatched_fields = []
        for col in common_cols:
            col_a, col_b = f"{col}_a", f"{col}_b"
            val_a = row[col_a] if col_a in row else row[col]
            val_b = row[col_b] if col_b in row else row[col]
            if not values_equal(val_a, val_b, tolerance):
                mismatched_fields.append((col, val_a, val_b))

        if mismatched_fields:
            for col, val_a, val_b in mismatched_fields:
                mismatch_rows.append({
                    "key": key_val,
                    "field": col,
                    "value_a": val_a,
                    "value_b": val_b,
                })
        else:
            record = {"key": key_val}
            for col in common_cols:
                col_a = f"{col}_a"
                record[col] = row[col_a] if col_a in row else row[col]
            clean_rows.append(record)

    results = {
        "clean": pd.DataFrame(clean_rows),
        "mismatches": pd.DataFrame(mismatch_rows),
        "only_in_a": only_in_a,
        "only_in_b": only_in_b,
        "duplicates_a": duplicates_a,
        "duplicates_b": duplicates_b,
    }

    schema_notes = {
        "only_in_a_columns": only_in_a_cols,
        "only_in_b_columns": only_in_b_cols,
    }

    return results, schema_notes


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------
def print_summary(results, schema_notes):
    print("=" * 60)
    print("RECONCILIATION SUMMARY")
    print("=" * 60)
    print(f"Clean matches:        {len(results['clean'])}")
    print(f"Field mismatches:     {len(results['mismatches'])} field-level diffs")
    print(f"Only in A:            {len(results['only_in_a'])}")
    print(f"Only in B:            {len(results['only_in_b'])}")
    print(f"Duplicate keys in A:  {len(results['duplicates_a'])} rows")
    print(f"Duplicate keys in B:  {len(results['duplicates_b'])} rows")

    if schema_notes["only_in_a_columns"]:
        print(f"\nColumns only in file A (not compared): {schema_notes['only_in_a_columns']}")
    if schema_notes["only_in_b_columns"]:
        print(f"Columns only in file B (not compared): {schema_notes['only_in_b_columns']}")

    if len(results["mismatches"]):
        print("\n--- Mismatches ---")
        print(results["mismatches"].to_string(index=False))

    if len(results["duplicates_a"]) or len(results["duplicates_b"]):
        print("\n--- Duplicate keys (excluded from pairing) ---")
        if len(results["duplicates_a"]):
            print("File A:")
            print(results["duplicates_a"].to_string(index=False))
        if len(results["duplicates_b"]):
            print("File B:")
            print(results["duplicates_b"].to_string(index=False))
    print("=" * 60)


def write_outputs(results, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    for name, df in results.items():
        df.to_csv(os.path.join(output_dir, f"{name}.csv"), index=False)
    print(f"\nCSV reports written to: {output_dir}/")


# ---------------------------------------------------------------------------
# Excel report (formatted, multi-tab)
# ---------------------------------------------------------------------------
_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_MISMATCH_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_ONLY_A_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_ONLY_B_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
_DUPLICATE_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
_THIN_SIDE = Side(style="thin", color="D9D9D9")
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def _write_sheet(wb, title, df, row_fill=None, cell_fills=None):
    """Write one DataFrame as a styled sheet.

    row_fill highlights every data row (used for buckets where the whole row
    needs attention, like duplicates or orphaned keys). cell_fills highlights
    only specific columns by name (used for mismatches, to flag the two
    differing values rather than the whole row).
    """
    ws = wb.create_sheet(title)
    columns = list(df.columns)

    if not columns:
        ws.append(["(no rows)"])
        return ws

    ws.append(columns)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    if df.empty:
        ws.append(["(no rows)"] + [""] * (len(columns) - 1))
    else:
        for row in df.itertuples(index=False):
            ws.append(list(row))

        col_idx = {name: i + 1 for i, name in enumerate(columns)}
        for r in range(2, ws.max_row + 1):
            if row_fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = row_fill
            if cell_fills:
                for col_name, fill in cell_fills.items():
                    idx = col_idx.get(col_name)
                    if idx:
                        ws.cell(row=r, column=idx).fill = fill

    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).border = _THIN_BORDER

    for c in range(1, ws.max_column + 1):
        letter = get_column_letter(c)
        max_len = max(
            (len(str(ws.cell(row=r, column=c).value)) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 50)

    return ws


def _write_summary(wb, results, schema_notes):
    ws = wb.create_sheet("Summary", 0)
    ws.append(["Reconciliation Summary"])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(["Bucket", "Count"])
    for cell in ws[header_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    for label, count in [
        ("Clean matches", len(results["clean"])),
        ("Field-level mismatches", len(results["mismatches"])),
        ("Only in A", len(results["only_in_a"])),
        ("Only in B", len(results["only_in_b"])),
        ("Duplicate keys in A", len(results["duplicates_a"])),
        ("Duplicate keys in B", len(results["duplicates_b"])),
    ]:
        ws.append([label, count])

    if schema_notes["only_in_a_columns"] or schema_notes["only_in_b_columns"]:
        ws.append([])
        ws.append(["Schema notes (columns not compared)"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        if schema_notes["only_in_a_columns"]:
            ws.append(["Columns only in file A", ", ".join(schema_notes["only_in_a_columns"])])
        if schema_notes["only_in_b_columns"]:
            ws.append(["Columns only in file B", ", ".join(schema_notes["only_in_b_columns"])])

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 45
    return ws


def write_excel_report(results, schema_notes, path):
    wb = Workbook()
    wb.remove(wb.active)

    _write_summary(wb, results, schema_notes)
    _write_sheet(wb, "Clean", results["clean"])
    _write_sheet(
        wb, "Mismatches", results["mismatches"],
        cell_fills={"value_a": _MISMATCH_FILL, "value_b": _MISMATCH_FILL},
    )
    _write_sheet(wb, "Only in A", results["only_in_a"], row_fill=_ONLY_A_FILL)
    _write_sheet(wb, "Only in B", results["only_in_b"], row_fill=_ONLY_B_FILL)
    _write_sheet(wb, "Duplicates A", results["duplicates_a"], row_fill=_DUPLICATE_FILL)
    _write_sheet(wb, "Duplicates B", results["duplicates_b"], row_fill=_DUPLICATE_FILL)

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    print(f"Excel report written to: {path}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def parse_column_mapping(raw):
    """Parse 'file_b_col=file_a_col,other_b=other_a' into a dict."""
    mapping = {}
    for pair in raw.split(","):
        pair = pair.strip()
        if not pair:
            continue
        if "=" not in pair:
            raise argparse.ArgumentTypeError(
                f"Invalid --column-mapping entry '{pair}' (expected format: file_b_col=file_a_col)"
            )
        b_col, a_col = pair.split("=", 1)
        mapping[b_col.strip()] = a_col.strip()
    return mapping


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Reconcile two spreadsheets (CSV or XLSX) by a key column. "
                     "Any flag left unset falls back to the CONFIG dict in main.py.",
    )
    parser.add_argument("--file-a", default=CONFIG["file_a"], help="Path to file A")
    parser.add_argument("--file-b", default=CONFIG["file_b"], help="Path to file B")
    parser.add_argument("--key-column", default=CONFIG["key_column"], help="Key column name")
    parser.add_argument(
        "--tolerance", type=float, default=CONFIG["numeric_tolerance"],
        help="Numeric comparison tolerance",
    )
    parser.add_argument(
        "--column-mapping", type=parse_column_mapping, default=None,
        help="Comma-separated file_b_col=file_a_col pairs, e.g. 'full_name=name,amt=amount'",
    )
    parser.add_argument("--output-dir", default=CONFIG["output_dir"], help="Directory for CSV output")
    parser.add_argument("--output-excel", default=CONFIG["output_excel_path"], help="Path for the Excel report")
    return parser.parse_args(argv)


def build_config(args):
    config = dict(CONFIG)
    config["file_a"] = args.file_a
    config["file_b"] = args.file_b
    config["key_column"] = args.key_column
    config["numeric_tolerance"] = args.tolerance
    config["output_dir"] = args.output_dir
    config["output_excel_path"] = args.output_excel
    if args.column_mapping is not None:
        config["column_mapping"] = args.column_mapping
    return config


def main():
    config = build_config(parse_args())

    try:
        results, schema_notes = reconcile(config)
    except ReconciliationError as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

    print_summary(results, schema_notes)
    write_outputs(results, config["output_dir"])
    write_excel_report(results, schema_notes, config["output_excel_path"])


if __name__ == "__main__":
    main()
